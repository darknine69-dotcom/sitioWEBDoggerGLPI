"""Prueba el módulo del Observador.

Se fija sobre todo el caso que se rompió: al agregar una persona a la lista
personal se guardaba sin usuario y sin correo, así que el observador perdía
de vista a quién había agregado y la fila terminaba contando todos los
tickets de sus zonas.
"""
from io import BytesIO

from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from apps.accounts.models import Usuario
from apps.tickets.models import (
    ElementoMiLista,
    MonitoreoAuditoria,
    ObservadorPunto,
    Punto,
    Ticket,
    ZonaObservador,
)


@override_settings(STORAGES={
    "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
    "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
})
class ObservadorListaTest(TestCase):
    def setUp(self):
        self.observador = Usuario.objects.create_user(
            email="obs@x.com", password="x", nombre="Obs", rol=Usuario.Rol.OBSERVADOR
        )
        self.otro = Usuario.objects.create_user(
            email="otro@x.com", password="x", nombre="Otro", rol=Usuario.Rol.OBSERVADOR
        )
        self.zona = Punto.objects.create(nombre="Zona Uno")
        self.ajena = Punto.objects.create(nombre="Zona Ajena")
        ObservadorPunto.objects.create(observador=self.observador, punto=self.zona)
        self.client.force_login(self.observador)

    def _ticket(self, nombre, punto, email="", estado=Ticket.Estado.ABIERTO,
                prioridad=Ticket.Prioridad.MEDIA):
        return Ticket.objects.create(
            titulo="T",
            descripcion="d",
            prioridad=prioridad,
            estado=estado,
            solicitante_nombre=nombre,
            solicitante_email=email or None,
            solicitante_punto=punto,
        )

    def _agregar(self, correo):
        return self.client.post(
            reverse("tickets:obs_agregar"),
            {"email": correo, "volver": reverse("tickets:obs_mi_lista")},
        )

    # --- el bug: la persona se perdia al agregarla ---
    def test_agregar_persona_sin_usuario_guarda_el_correo(self):
        self._agregar("afectado@x.com")
        fila = ElementoMiLista.objects.get(observador=self.observador)
        self.assertEqual(fila.usuario_email, "afectado@x.com")
        self.assertEqual(fila.tipo, ElementoMiLista.TIPO_USUARIO)

    def test_agregar_persona_con_usuario_lo_enlaza(self):
        registrado = Usuario.objects.create_user(
            email="conuser@x.com", password="x", nombre="Con User",
            rol=Usuario.Rol.USUARIO,
        )
        self._agregar("conuser@x.com")
        fila = ElementoMiLista.objects.get(observador=self.observador)
        self.assertEqual(fila.usuario_email, "conuser@x.com")
        self.assertEqual(fila.usuario_id, registrado.pk)

    def test_agregar_dos_veces_no_duplica(self):
        self._agregar("afectado@x.com")
        self._agregar("afectado@x.com")
        self.assertEqual(
            ElementoMiLista.objects.filter(observador=self.observador).count(), 1
        )

    def test_mi_lista_no_falla_con_persona_sin_usuario(self):
        """Antes esto reventaba con AttributeError al leer usuario_email."""
        self._agregar("afectado@x.com")
        for url in ("tickets:obs_mi_lista", "tickets:obs_monitoreo", "tickets:obs_usuarios"):
            with self.subTest(url=url):
                self.assertEqual(self.client.get(reverse(url)).status_code, 200)

    def test_mi_lista_solo_cuenta_los_tickets_de_esa_persona(self):
        self._ticket("Afectado", "Zona Uno", email="afectado@x.com")
        self._ticket("Afectado", "Zona Uno", email="afectado@x.com")
        self._ticket("Otra", "Zona Uno", email="otra@x.com")
        self._agregar("afectado@x.com")
        respuesta = self.client.get(reverse("tickets:obs_mi_lista"))
        self.assertEqual(respuesta.context["filas"][0]["total"], 2)
        self.assertEqual(respuesta.context["filas"][0]["nombre"], "afectado@x.com")

    def test_mi_lista_es_solo_el_panel_de_zonas(self):
        """Mi lista ya no tiene tabla: solo el panel de dos columnas.

        Y ese panel ya no lleva título, ni «Mis puntos», ni «Sin zona»:
        lo único que se pinta son las zonas.
        """
        ZonaObservador.objects.create(observador=self.observador, nombre="Norte")
        self._agregar("sin-zona@x.com")
        respuesta = self.client.get(reverse("tickets:obs_mi_lista"))
        html = respuesta.content.decode()
        # El panel sí está, con sus zonas.
        self.assertIn('class="pan-grid"', html)
        self.assertIn("Mis zonas", html)
        self.assertIn("Supervisor a cargo", html)
        # La tabla de antes ya no existe.
        self.assertNotIn("Técnico", html)
        # Ni el título, ni los dos grupos que se quitaron.
        self.assertNotIn("Panel corporativo", html)
        self.assertNotIn("Mis puntos", html)
        self.assertNotIn("Sin zona", html)

    def test_exportar_no_falla_con_persona_sin_usuario(self):
        self._agregar("afectado@x.com")
        self.assertEqual(self.client.get(reverse("tickets:obs_exportar")).status_code, 200)

    # --- aislamiento por zona ---
    def test_no_ve_personas_de_zonas_ajenas(self):
        self._ticket("De Mi Zona", "Zona Uno", email="mio@x.com")
        self._ticket("De Zona Ajena", "Zona Ajena", email="ajeno@x.com")
        respuesta = self.client.get(reverse("tickets:obs_usuarios"))
        self.assertContains(respuesta, "mio@x.com")
        self.assertNotContains(respuesta, "ajeno@x.com")

    def test_mi_lista_no_arrastra_tickets_de_zona_ajena(self):
        self._ticket("Afectado", "Zona Uno", email="afectado@x.com")
        self._ticket("Afectado", "Zona Ajena", email="afectado@x.com")
        self._agregar("afectado@x.com")
        respuesta = self.client.get(reverse("tickets:obs_mi_lista"))
        self.assertEqual(respuesta.context["filas"][0]["total"], 1)

    def test_otro_observador_no_roba_elementos(self):
        self._agregar("afectado@x.com")
        self.client.force_login(self.otro)
        respuesta = self.client.get(reverse("tickets:obs_mi_lista"))
        self.assertEqual(respuesta.context["filas"], [])

    # --- filtros de la vista de usuarios ---
    def test_filtro_por_texto(self):
        self._ticket("Ana Ruiz", "Zona Uno", email="ana@x.com")
        self._ticket("Beto Diaz", "Zona Uno", email="beto@x.com")
        respuesta = self.client.get(reverse("tickets:obs_usuarios"), {"q": "ana"})
        self.assertContains(respuesta, "ana@x.com")
        self.assertNotContains(respuesta, "beto@x.com")

    def test_filtro_por_zona(self):
        ObservadorPunto.objects.create(observador=self.observador, punto=self.ajena)
        self._ticket("Ana", "Zona Uno", email="ana@x.com")
        self._ticket("Beto", "Zona Ajena", email="beto@x.com")
        respuesta = self.client.get(
            reverse("tickets:obs_usuarios"), {"punto": "Zona Ajena"}
        )
        self.assertContains(respuesta, "beto@x.com")
        self.assertNotContains(respuesta, "ana@x.com")

    def test_filtro_activos_exige_ticket_abierto(self):
        self._ticket("Ana", "Zona Uno", email="ana@x.com")
        self._ticket("Beto", "Zona Uno", email="beto@x.com",
                     estado=Ticket.Estado.CERRADO)
        respuesta = self.client.get(reverse("tickets:obs_usuarios"), {"estado": "activos"})
        self.assertContains(respuesta, "ana@x.com")
        self.assertNotContains(respuesta, "beto@x.com")

    def test_filtro_alerta_exige_urgente_y_abierto(self):
        self._ticket("Ana", "Zona Uno", email="ana@x.com",
                     prioridad=Ticket.Prioridad.URGENTE)
        self._ticket("Beto", "Zona Uno", email="beto@x.com",
                     estado=Ticket.Estado.CERRADO,
                     prioridad=Ticket.Prioridad.URGENTE)
        respuesta = self.client.get(reverse("tickets:obs_usuarios"), {"estado": "alerta"})
        self.assertContains(respuesta, "ana@x.com")
        self.assertNotContains(respuesta, "beto@x.com")

    def test_orden_por_nombre(self):
        self._ticket("Zulma", "Zona Uno", email="zulma@x.com")
        self._ticket("Ana", "Zona Uno", email="ana@x.com")
        respuesta = self.client.get(reverse("tickets:obs_usuarios"), {"orden": "nombre"})
        # La tabla es única y también trae a las cuentas (Obs, Otro), así que
        # se compara el orden de la gente con tickets.
        nombres = [
            f["nombre"] for f in respuesta.context["filas"] if f["total"] > 0
        ]
        self.assertEqual(nombres, ["Ana", "Zulma"])

    def test_orden_por_abiertos(self):
        self._ticket("Ana", "Zona Uno", email="ana@x.com")
        self._ticket("Beto", "Zona Uno", email="beto@x.com")
        self._ticket("Beto", "Zona Uno", email="beto@x.com")
        respuesta = self.client.get(reverse("tickets:obs_usuarios"), {"orden": "abiertos"})
        nombres = [
            f["nombre"] for f in respuesta.context["filas"] if f["total"] > 0
        ]
        self.assertEqual(nombres, ["Beto", "Ana"])

    def test_usuarios_muestra_el_total_sin_filtro(self):
        self._ticket("Ana", "Zona Uno", email="ana@x.com")
        self._ticket("Beto", "Zona Uno", email="beto@x.com")
        respuesta = self.client.get(reverse("tickets:obs_usuarios"), {"q": "ana"})
        # La lista única junta las dos personas con tickets con las cuentas
        # del observador (Obs y Otro), y el filtro aprieta sobre las cuatro.
        self.assertEqual(respuesta.context["total"], 1)
        self.assertEqual(respuesta.context["total_sin_filtro"], 4)

    def test_usuarios_trae_tabla_total_para_meter_a_zona(self):
        Usuario.objects.create_user(
            email="nueva@x.com", password="x", nombre="Nueva",
            rol=Usuario.Rol.USUARIO, ubicacion="Zona Uno",
        )
        respuesta = self.client.get(reverse("tickets:obs_usuarios"))
        self.assertEqual(respuesta.status_code, 200)
        todos = respuesta.context["todos_usuarios"]
        self.assertTrue(any(u["email"] == "nueva@x.com" for u in todos))
        html = respuesta.content.decode()
        self.assertIn("Todos los usuarios", html)
        self.assertIn("Agregar", html)
        self.client.post(reverse("tickets:obs_agregar"), {"email": "nueva@x.com"})
        html = self.client.get(reverse("tickets:obs_usuarios")).content.decode()
        # La zona se pone en la propia columna, con el selector en línea.
        self.assertIn('class="obs-zona-form"', html)
        self.assertIn("Zona de Nueva", html)

    def test_panel_es_tablero_con_indicadores_propios(self):
        import re

        self._ticket("Ana", "Zona Uno", email="ana@x.com")
        self._agregar("ana@x.com")
        respuesta = self.client.get(reverse("tickets:obs_panel"))
        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta.context["total_personas"], 1)
        self.assertEqual(respuesta.context["total_tickets"], 1)
        self.assertEqual(respuesta.context["abiertos"], 1)
        self.assertEqual(respuesta.context["mi_lista_total"], 1)
        clases = set(re.findall(r"dash-hov-[a-z]+", respuesta.content.decode()))
        self.assertEqual(
            clases,
            {
                "dash-hov-usuarios", "dash-hov-total", "dash-hov-abiertos",
                "dash-hov-vencidos", "dash-hov-zonas", "dash-hov-progreso",
            },
        )

    def test_panel_sin_duplicados_y_con_buscador(self):
        html = self.client.get(reverse("tickets:obs_panel")).content.decode()
        self.assertNotIn("obs_usuarios", html)
        self.assertNotIn("Imprimir", html)
        self.assertNotIn("Exportar", html)
        self.assertIn("Buscar y agregar", html)
        self.assertIn("axFab", html)

    def test_buscador_del_panel_agrega_a_la_lista(self):
        Usuario.objects.create_user(
            email="nueva@x.com", password="x", nombre="Nueva",
            rol=Usuario.Rol.USUARIO,
        )
        respuesta = self.client.get(
            reverse("tickets:obs_panel"), {"q_busc": "nueva"}
        )
        self.assertContains(respuesta, "nueva@x.com")
        self.client.post(
            reverse("tickets:obs_agregar"), {"email": "nueva@x.com"}
        )
        self.assertTrue(
            ElementoMiLista.objects.filter(
                observador=self.observador, usuario_email="nueva@x.com"
            ).exists()
        )

    def test_cuenta_creada_se_ve_aunque_no_tenga_tickets(self):
        Usuario.objects.create_user(
            email="nueva@x.com", password="x", nombre="Nueva",
            rol=Usuario.Rol.USUARIO, ubicacion="Zona Uno",
        )
        respuesta = self.client.get(reverse("tickets:obs_panel"))
        filas = respuesta.context["personas"]
        nueva = [f for f in filas if f["email"] == "nueva@x.com"]
        self.assertEqual(len(nueva), 1)
        self.assertEqual(nueva[0]["total"], 0)
        self.assertIn("sin tickets aún", respuesta.content.decode())

    def test_panel_muestra_el_tecnico_de_cada_cuenta(self):
        tec = Usuario.objects.create_user(
            email="tec@x.com", password="x", nombre="Tec",
            rol=Usuario.Rol.TECNICO,
        )
        Ticket.objects.create(
            titulo="T", descripcion="d", solicitante_nombre="Ana",
            solicitante_email="ana@x.com", solicitante_punto="Zona Uno",
            tecnico_asignado=tec,
        )
        respuesta = self.client.get(reverse("tickets:obs_panel"))
        ana = [
            f for f in respuesta.context["personas"] if f["email"] == "ana@x.com"
        ]
        self.assertEqual(ana[0]["tecnico"], "Tec")
        self.assertIn("Tec", respuesta.content.decode())

    def test_panel_es_el_monitoreo_por_zonas(self):
        from apps.tickets.models import ZonaObservador

        ticket = self._ticket("Ana", "Zona Uno", email="ana@x.com")
        self._agregar("ana@x.com")
        zona = ZonaObservador.objects.create(observador=self.observador, nombre="Norte")
        elemento = ElementoMiLista.objects.get(observador=self.observador)
        elemento.zona = zona
        elemento.save()
        respuesta = self.client.get(reverse("tickets:obs_panel"))
        self.assertEqual(len(respuesta.context["grupos"]), 1)
        html = respuesta.content.decode()
        self.assertIn("Monitoreo por zonas", html)
        self.assertIn("Norte", html)
        self.assertIn("Técnico a cargo", html)
        self.assertIn(ticket.codigo, html)

    def test_menu_lateral_sin_panel_y_monitoreo_primero(self):
        html = self.client.get(reverse("tickets:obs_monitoreo")).content.decode()
        self.assertNotIn("Puntos asignados", html)
        self.assertNotIn('sidebar-label">Panel', html)
        self.assertIn('sidebar-label">Monitoreo', html)
        self.assertNotIn('sidebar-label">Buscar y agregar', html)

    def test_usuarios_y_lista_sin_botones_duplicados(self):
        html_usuarios = self.client.get(
            reverse("tickets:obs_usuarios")
        ).content.decode()
        self.assertNotIn("Imprimir", html_usuarios)
        self.assertNotIn("obs_buscar", html_usuarios)
        html_lista = self.client.get(
            reverse("tickets:obs_mi_lista")
        ).content.decode()
        self.assertNotIn("Agregar más", html_lista)

    def test_entrada_del_observador_es_monitoreo(self):
        from apps.accounts.routing import landing_por_rol

        self.assertEqual(
            landing_por_rol(self.observador),
            reverse("tickets:obs_monitoreo"),
        )


@override_settings(STORAGES={
    "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
    "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
})
class ObservadorAuditoriaMiListaTest(TestCase):
    """Mi lista trae abajo el historial de auditoría: la misma tabla que
    tiene el administrador, con sus contadores, sus columnas y sus filtros."""

    def setUp(self):
        self.observador = Usuario.objects.create_user(
            email="obs@x.com", password="x", nombre="Obs",
            rol=Usuario.Rol.OBSERVADOR,
        )
        self.otro = Usuario.objects.create_user(
            email="obs2@x.com", password="x", nombre="Otro",
            rol=Usuario.Rol.OBSERVADOR,
        )
        self.client.force_login(self.observador)
        MonitoreoAuditoria.objects.create(
            actor=self.observador, observador=self.observador,
            accion=MonitoreoAuditoria.CREAR_ZONA, zona_nombre="Norte",
            detalle="crea la zona Norte",
        )

    def test_mi_lista_trae_el_historial_con_su_tabla(self):
        html = self.client.get(reverse("tickets:obs_mi_lista")).content.decode()
        self.assertIn('id="auditoria"', html)
        self.assertIn("Registro de auditoría", html)
        # Las mismas columnas que la tabla del administrador.
        for columna in ("ID", "Fecha", "Hace", "Acción", "Zona", "Supervisor",
                        "Cuenta", "Hecho por", "Detalle", "Motivo", "Resultado",
                        "Versión"):
            with self.subTest(columna=columna):
                self.assertIn(f"<th>{columna}</th>", html)
        self.assertIn("crea la zona Norte", html)

    def test_mi_lista_trae_los_filtros_del_historial(self):
        html = self.client.get(reverse("tickets:obs_mi_lista")).content.decode()
        for control in ('id="aud-buscar"', 'id="aud-f-desde"', 'id="aud-f-hasta"',
                        'id="aud-f-accion"', 'id="aud-f-resultado"',
                        'id="aud-tabla"'):
            with self.subTest(control=control):
                self.assertIn(control, html)

    def test_los_contadores_de_hoy_y_de_la_semana_van_en_el_contexto(self):
        respuesta = self.client.get(reverse("tickets:obs_mi_lista"))
        self.assertEqual(respuesta.context["aud_hoy"], 1)
        self.assertEqual(respuesta.context["aud_semana"], 1)

    def test_muestra_todos_los_registros_igual_que_el_administrador(self):
        """A un observador le sale también lo que hicieron los demás."""
        MonitoreoAuditoria.objects.create(
            actor=self.otro, observador=self.otro,
            accion=MonitoreoAuditoria.ELIMINAR_ZONA, zona_nombre="Sur",
            detalle="borra la zona Sur",
        )
        html = self.client.get(reverse("tickets:obs_mi_lista")).content.decode()
        self.assertIn("crea la zona Norte", html)
        self.assertIn("borra la zona Sur", html)

    def test_el_historial_no_ofrece_la_exportacion_del_administrador(self):
        """Lo único que no puede bajar es el Excel completo del
        administrador: el suyo es otro, con solo sus zonas, y solo si la
        zona tiene el permiso (lo prueba ObservadorExportarAuditoriaTest)."""
        html = self.client.get(reverse("tickets:obs_mi_lista")).content.decode()
        self.assertNotIn(reverse("tickets:monitoreo_auditoria_exportar"), html)


@override_settings(STORAGES={
    "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
    "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
})
class ObservadorPermisosTest(TestCase):
    def test_otro_rol_no_entra(self):
        usuario = Usuario.objects.create_user(
            email="admin@x.com", password="x", nombre="A", rol=Usuario.Rol.ADMIN
        )
        self.client.force_login(usuario)
        for nombre in ("obs_panel", "obs_usuarios", "obs_buscar", "obs_mi_lista", "obs_monitoreo"):
            with self.subTest(nombre=nombre):
                self.assertEqual(self.client.get(reverse("tickets:" + nombre)).status_code, 403)


@override_settings(STORAGES={
    "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
    "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
})
class ObservadorBuscadorTest(TestCase):
    def setUp(self):
        self.observador = Usuario.objects.create_user(
            email="obs@x.com", password="x", nombre="Obs", rol=Usuario.Rol.OBSERVADOR
        )
        ObservadorPunto.objects.create(
            observador=self.observador, punto=Punto.objects.create(nombre="Zona Uno")
        )
        self.client.force_login(self.observador)

    def _ticket(self, nombre, punto, email=""):
        return Ticket.objects.create(
            titulo="T", descripcion="d",
            prioridad=Ticket.Prioridad.MEDIA,
            estado=Ticket.Estado.ABIERTO,
            solicitante_nombre=nombre,
            solicitante_email=email or None,
            solicitante_punto=punto,
        )

    def test_busca_por_nombre(self):
        self._ticket("Ana Ruiz", "Zona Uno", email="ana@x.com")
        respuesta = self.client.get(reverse("tickets:obs_buscar"), {"q": "ana"})
        self.assertContains(respuesta, "ana@x.com")

    def test_no_busca_fuera_de_sus_zonas(self):
        self._ticket("Ajena", "Zona Ajena", email="ajena@x.com")
        respuesta = self.client.get(reverse("tickets:obs_buscar"), {"q": "ajena"})
        self.assertNotContains(respuesta, "ajena@x.com")

    def test_filtro_tipo_punto(self):
        self._ticket("Ana", "Zona Uno", email="ana@x.com")
        respuesta = self.client.get(
            reverse("tickets:obs_buscar"), {"q": "Zona", "tipo": "punto"}
        )
        personas = respuesta.context["resultados"]["personas"]
        puntos = respuesta.context["resultados"]["puntos"]
        self.assertEqual(personas, [])
        self.assertEqual([p["nombre"] for p in puntos], ["Zona Uno"])

    def test_filtro_tipo_persona(self):
        self._ticket("Ana", "Zona Uno", email="ana@x.com")
        respuesta = self.client.get(
            reverse("tickets:obs_buscar"), {"q": "Ana", "tipo": "persona"}
        )
        self.assertEqual(
            [p["nombre"] for p in respuesta.context["resultados"]["personas"]], ["Ana"]
        )
        self.assertEqual(respuesta.context["resultados"]["puntos"], [])

    def test_marca_que_ya_esta_en_la_lista(self):
        self._ticket("Ana", "Zona Uno", email="ana@x.com")
        self.client.post(
            reverse("tickets:obs_agregar"),
            {"email": "ana@x.com", "volver": reverse("tickets:obs_buscar")},
        )
        respuesta = self.client.get(reverse("tickets:obs_buscar"), {"q": "ana"})
        self.assertTrue(respuesta.context["resultados"]["personas"][0]["ya_agregado"])


class ObservadorCrearZonaTest(TestCase):
    """El observador crea sus zonas desde Usuarios, con la misma ventana
    que la del administrador: validación en vivo, pestaña de permisos y
    avisos por toast al crear o al fallar."""

    def setUp(self):
        self.observador = Usuario.objects.create_user(
            email="obs@x.com", password="x", nombre="Obs",
            rol=Usuario.Rol.OBSERVADOR,
        )
        self.client.force_login(self.observador)

    def _crear(self, follow=False, **datos):
        base = {"nombre": "Sede norte", "volver": reverse("tickets:obs_usuarios")}
        base.update(datos)
        return self.client.post(reverse("tickets:obs_zona_crear"), base, follow=follow)

    def test_usuarios_trae_el_boton_y_la_ventana_de_crear_zona(self):
        html = self.client.get(reverse("tickets:obs_usuarios")).content.decode()
        self.assertIn('id="cz-abrir"', html)
        self.assertIn("Crear zona", html)
        self.assertIn('id="cz-modal"', html)
        self.assertIn(reverse("tickets:obs_zona_crear"), html)
        # Las notificaciones y las ventanas emergentes de la zona.
        self.assertIn('name="notif_correo"', html)
        self.assertIn('name="notif_popup"', html)
        self.assertIn('name="alertas_ans"', html)
        self.assertIn("data-cz-tab", html)

    def test_crear_zona_la_deja_a_su_nombre_y_abre_la_tabla(self):
        respuesta = self._crear()
        zona = ZonaObservador.objects.get(observador=self.observador)
        self.assertEqual(zona.nombre, "Sede norte")
        # Vuelve a la tabla de Todos los usuarios, señalando la zona nueva.
        self.assertRedirects(
            respuesta,
            f"{reverse('tickets:obs_usuarios')}?zona_nueva={zona.pk}",
        )

    def test_recien_creada_propone_meterle_puntos_de_venta(self):
        """La ventana de la zona nueva ofrece pasar a la tabla para
        asignarle puntos de venta."""
        self._crear(nombre="Sede norte")
        zona = ZonaObservador.objects.get(observador=self.observador)
        html = self.client.get(
            reverse("tickets:obs_usuarios"), {"zona_nueva": zona.pk}
        ).content.decode()
        self.assertIn('id="zona-nueva"', html)
        self.assertIn("Zona «Sede norte» creada", html)
        # La aceptación lleva a la tabla filtrada por los puntos de venta
        # que aún no están en la lista.
        self.assertIn("rol=usuario", html)
        self.assertIn("estado=pendientes", html)

    def test_la_ventana_solo_sale_por_su_zona(self):
        """Nadie señala una zona que no sea del observador."""
        ajena = Usuario.objects.create_user(
            email="obs2@x.com", password="x", nombre="Otro",
            rol=Usuario.Rol.OBSERVADOR,
        )
        ajena_zona = ZonaObservador.objects.create(observador=ajena, nombre="Ajena")
        html = self.client.get(
            reverse("tickets:obs_usuarios"), {"zona_nueva": ajena_zona.pk}
        ).content.decode()
        self.assertNotIn('id="zona-nueva"', html)

    def test_zona_creada_sale_en_el_menu_de_zona_de_la_tabla(self):
        """La zona creada aparece en el filtro de zona de arriba —que antes
        salía vacío— y en el selector en línea de la columna Zona."""
        Usuario.objects.create_user(
            email="ana@x.com", password="x", nombre="Ana",
            rol=Usuario.Rol.USUARIO,
        )
        self.client.post(reverse("tickets:obs_agregar"), {"email": "ana@x.com"})
        self._crear(nombre="Sede norte")
        zona = ZonaObservador.objects.get(observador=self.observador)
        html = self.client.get(reverse("tickets:obs_usuarios")).content.decode()
        self.assertIn(f'<option value="{zona.nombre}">{zona.nombre}</option>', html)
        self.assertIn(f'<option value="{zona.pk}">{zona.nombre}</option>', html)
        self.assertIn('class="obs-zona-form"', html)

    def test_zona_creada_se_ve_en_monitoreo_estando_vacia(self):
        """Una zona nueva aparece en el monitoreo aunque aún no tenga a
        nadie, igual que en el tablero del administrador."""
        self._crear(nombre="Sede norte")
        respuesta = self.client.get(reverse("tickets:obs_monitoreo"))
        self.assertContains(respuesta, "Sede norte")
        self.assertEqual(respuesta.context["total_zonas"], 1)

    def test_el_filtro_de_zona_encuentra_a_la_gente_de_esa_zona(self):
        """El desplegable de zona filtra de verdad la tabla."""
        self._crear(nombre="Sede norte")
        zona = ZonaObservador.objects.get(observador=self.observador)
        Usuario.objects.create_user(
            email="ana@x.com", password="x", nombre="Ana", rol=Usuario.Rol.USUARIO
        )
        # Ana está en la lista y ya en la zona; Beto se queda sin ella.
        self.client.post(reverse("tickets:obs_agregar"), {"email": "ana@x.com"})
        self.client.post(reverse("tickets:obs_agregar"), {"email": "beto@x.com"})
        ana = ElementoMiLista.objects.get(
            observador=self.observador, usuario_email="ana@x.com"
        )
        ana.zona = zona
        ana.save()
        respuesta = self.client.get(
            reverse("tickets:obs_usuarios"), {"punto": "Sede norte"}
        )
        self.assertContains(respuesta, "ana@x.com")
        self.assertNotContains(respuesta, "beto@x.com")

    def test_crear_zona_sin_nombre_no_crea_y_avisa(self):
        respuesta = self.client.post(
            reverse("tickets:obs_zona_crear"),
            {"nombre": "   ", "volver": reverse("tickets:obs_usuarios")},
            follow=True,
        )
        self.assertFalse(ZonaObservador.objects.exists())
        mensajes = [m.message for m in respuesta.context["messages"]]
        self.assertTrue(any("nombre" in m.lower() for m in mensajes))

    def test_crear_zona_repite_nombre_no_duplica_y_avisa(self):
        ZonaObservador.objects.create(observador=self.observador, nombre="Sede norte")
        respuesta = self._crear(nombre="SEDE NORTE", follow=True)
        self.assertEqual(ZonaObservador.objects.filter(observador=self.observador).count(), 1)
        mensajes = [m.message for m in respuesta.context["messages"]]
        self.assertTrue(any("SEDE NORTE" in m for m in mensajes))

    def test_crear_zona_guarda_las_notificaciones_del_formulario(self):
        # La casilla apagada no viaja en el POST: hay que guardarla apagada.
        self._crear(nombre="Sin avisos", notif_correo="on")
        zona = ZonaObservador.objects.get(nombre="Sin avisos")
        self.assertTrue(zona.notif_correo)
        self.assertFalse(zona.notif_popup)

    def test_crear_zona_sin_permiso_de_exportar_no_lo_deja_marcado(self):
        self._crear(nombre="Sin exportar")
        self.assertFalse(ZonaObservador.objects.get(nombre="Sin exportar").perm_exportar)

    def test_crear_zona_deja_auditoria(self):
        self._crear(nombre="Con registro")
        self.assertTrue(
            MonitoreoAuditoria.objects.filter(
                observador=self.observador,
                accion=MonitoreoAuditoria.CREAR_ZONA,
                zona_nombre="Con registro",
            ).exists()
        )

    def test_crear_zona_de_este_observador_no_toca_las_de_los_demas(self):
        otro = Usuario.objects.create_user(
            email="otro@x.com", password="x", nombre="Otro",
            rol=Usuario.Rol.OBSERVADOR,
        )
        ZonaObservador.objects.create(observador=otro, nombre="Ajena")
        self._crear(nombre="Mia")
        self.assertEqual(
            list(
                ZonaObservador.objects.filter(observador=otro)
                .values_list("nombre", flat=True)
            ),
            ["Ajena"],
        )

    def test_zona_creada_desde_usuarios_se_ve_en_zonas(self):
        self._crear(nombre="Sede norte")
        html = self.client.get(reverse("tickets:obs_zonas")).content.decode()
        self.assertIn("Sede norte", html)


class ObservadorTablaUnicaDatosTest(TestCase):
    """La tabla única de Usuarios tiene que pintar los datos que ya están
    en la base: ubicación, técnico, abiertos, tickets y avance."""

    def setUp(self):
        self.observador = Usuario.objects.create_user(
            email="obs@x.com", password="x", nombre="Obs",
            rol=Usuario.Rol.OBSERVADOR,
        )
        self.zona = Punto.objects.create(nombre="Zona Uno")
        ObservadorPunto.objects.create(observador=self.observador, punto=self.zona)
        self.ajena = Punto.objects.create(nombre="Zona Ajena")
        self.client.force_login(self.observador)

    def _ticket(self, nombre, punto, email, estado=Ticket.Estado.ABIERTO):
        return Ticket.objects.create(
            titulo="T", descripcion="d", prioridad=Ticket.Prioridad.MEDIA,
            estado=estado, solicitante_nombre=nombre, solicitante_email=email,
            solicitante_punto=punto,
        )

    def _fila(self, respuesta, correo):
        filas = [f for f in respuesta.context["filas"] if f["email"] == correo]
        self.assertEqual(len(filas), 1, f"no hay fila para {correo}")
        return filas[0]

    def test_ubicacion_cae_al_punto_del_ticket_si_la_cuenta_no_tiene(self):
        """La cuenta sin ubicación no salía en blanco: se toma la del ticket."""
        Usuario.objects.create_user(
            email="ana@x.com", password="x", nombre="Ana",
            rol=Usuario.Rol.USUARIO, ubicacion="",
        )
        self._ticket("Ana", "Zona Uno", "ana@x.com")
        fila = self._fila(
            self.client.get(reverse("tickets:obs_usuarios")), "ana@x.com"
        )
        self.assertEqual(fila["ubicacion"], "Zona Uno")

    def test_cuentas_de_mi_lista_muestran_sus_datos_reales(self):
        """Una cuenta de mi lista cuyos tickets están fuera de mis puntos
        seguía saliendo en ceros: ahora se le ven sus números de verdad."""
        Usuario.objects.create_user(
            email="ana@x.com", password="x", nombre="Ana",
            rol=Usuario.Rol.USUARIO,
        )
        self._ticket("Ana", "Zona Ajena", "ana@x.com")
        self._ticket("Ana", "Zona Ajena", "ana@x.com", estado=Ticket.Estado.CERRADO)
        self.client.post(reverse("tickets:obs_agregar"), {"email": "ana@x.com"})

        fila = self._fila(
            self.client.get(reverse("tickets:obs_usuarios")), "ana@x.com"
        )
        self.assertEqual(fila["total"], 2)
        self.assertEqual(fila["abiertos"], 1)
        self.assertEqual(fila["cerrados_total"], 1)
        self.assertEqual(fila["avance"], 50)
        # La ubicación también sale del ticket, no de la cuenta.
        self.assertEqual(fila["ubicacion"], "Zona Ajena")

    def test_tabla_unica_incluye_a_quien_tiene_tickets_pero_no_cuenta(self):
        """Antes solo salía en la tabla filtrada, que ya no existe."""
        self._ticket("Sin Cuenta", "Zona Uno", "sin@x.com")
        respuesta = self.client.get(reverse("tickets:obs_usuarios"))
        fila = self._fila(respuesta, "sin@x.com")
        self.assertEqual(fila["rol"], "Sin cuenta")
        self.assertEqual(fila["total"], 1)
        self.assertEqual(fila["ubicacion"], "Zona Uno")

    def test_no_pinta_a_la_gente_de_las_zonas_ajenas(self):
        self._ticket("Ajeno", "Zona Ajena", "ajeno@x.com")
        html = self.client.get(reverse("tickets:obs_usuarios")).content.decode()
        self.assertNotIn("ajeno@x.com", html)

    def test_no_muestra_el_contador_de_tickets_fuera(self):
        """El «+N fuera» de la columna de abiertos ya no sale."""
        html = self.client.get(reverse("tickets:obs_usuarios")).content.decode()
        self.assertNotIn("fuera_puntos", html)
        self.assertNotIn("Tickets suyos fuera de tus puntos asignados", html)

    def test_los_filtros_van_justo_encima_de_la_tabla(self):
        """Los filtros van bajo los indicadores y pegados a la tabla, que
        es lo que filtran."""
        html = self.client.get(reverse("tickets:obs_usuarios")).content.decode()
        i_indicadores = html.index("dash-stat-row")
        i_filtros = html.index('class="obs-filtros obs-filtros-tabla"')
        i_tabla = html.index('id="tabla-usuarios"')
        self.assertLess(i_indicadores, i_filtros)
        self.assertLess(i_filtros, i_tabla)

    def test_la_busqueda_no_se_reemplaza_por_javascript(self):
        """El buscador manda el formulario como el resto de filtros: antes
        el JS cambiaba solo las filas y los indicadores se quedaban viejos."""
        html = self.client.get(reverse("tickets:obs_usuarios")).content.decode()
        self.assertNotIn("data-vivo", html)


@override_settings(STORAGES={
    "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
    "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
})
class ObservadorIndicadoresSonEnlacesTest(TestCase):
    """Los indicadores no son cajas mudas.

    Responden al paso del ratón (el CSS ya no distingue `<a>` de `<div>`)
    y, al pulsarlos, bajan dentro de la misma página a la sección que
    nombran. Un `<div>` no se puede pulsar, así que no debe quedar ninguno.
    """

    def setUp(self):
        self.observador = Usuario.objects.create_user(
            email="obs@x.com", password="x", nombre="Obs",
            rol=Usuario.Rol.OBSERVADOR,
        )
        self.client.force_login(self.observador)

    def test_los_del_monitoreo_llevan_a_su_seccion(self):
        html = self.client.get(reverse("tickets:obs_monitoreo")).content.decode()
        esperado = [
            ("dash-hov-zonas", "#zonas-monitoreo"),
            ("dash-hov-total", "#tickets-recientes"),
            ("dash-hov-abiertos", "#resumen-zonas"),
            ("dash-hov-ans", "#por-vencer"),
            ("dash-hov-resueltos", "#resumen-zonas"),
            ("dash-hov-usuarios", "#resumen-zonas"),
            ("dash-hov-vencidos", "#estado-lista"),
        ]
        for clase, destino in esperado:
            with self.subTest(clase=clase):
                # «Zonas a cargo» va primero y más grande, así que su etiqueta
                # lleva además `dash-stat--grande` dentro de la misma clase.
                self.assertIn(f'{clase}" href="{destino}"', html)
        # Cada destino existe de verdad, aunque no haya nada que enseñar:
        # si no, el enlace se queda sin sitio adónde ir.
        for ancla in ('id="zonas-monitoreo"', 'id="tickets-recientes"',
                      'id="resumen-zonas"', 'id="estado-lista"',
                      'id="por-vencer"'):
            with self.subTest(ancla=ancla):
                self.assertIn(ancla, html)

    def test_los_de_usuarios_llevan_a_lo_que_nombran(self):
        html = self.client.get(reverse("tickets:obs_usuarios")).content.decode()
        esperado = [
            ("dash-hov-usuarios", "#todos-usuarios"),
            ("dash-hov-total", "?orden=tickets"),
            ("dash-hov-abiertos", "?estado=activos"),
            ("dash-hov-resueltos", "?orden=avance"),
            ("dash-hov-progreso", "?estado=agregados"),
            ("dash-hov-vencidos", "?estado=alerta"),
        ]
        for clase, destino in esperado:
            with self.subTest(clase=clase):
                self.assertIn(f'class="dash-stat {clase}" href="{destino}"', html)
        self.assertIn('id="todos-usuarios"', html)

    def test_no_queda_ningun_indicador_como_caja_muda(self):
        for vista in ("obs_monitoreo", "obs_usuarios"):
            with self.subTest(vista=vista):
                html = self.client.get(reverse("tickets:" + vista)).content.decode()
                # La fila que los agrupa sí es un <div>, pero cada indicador
                # tiene que ser enlace: `<div class="dash-stat "`.
                self.assertNotIn('<div class="dash-stat ', html)


@override_settings(STORAGES={
    "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
    "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
})
class ObservadorExportarAuditoriaTest(TestCase):
    """El observador baja su auditoría a Excel.

    El permiso lo pone la zona, no el rol: si la carpeta se creó con la
    casilla «Exportar auditoría» marcada entra en el fichero; si no, ni
    botón ni descarga. Y de todas formas solo sale lo suyo.
    """

    def setUp(self):
        self.observador = Usuario.objects.create_user(
            email="obs@x.com", password="x", nombre="Obs",
            rol=Usuario.Rol.OBSERVADOR,
        )
        self.ajeno = Usuario.objects.create_user(
            email="otro@x.com", password="x", nombre="Otro",
            rol=Usuario.Rol.OBSERVADOR,
        )
        self.zona = ZonaObservador.objects.create(
            observador=self.observador, nombre="Norte"
        )
        MonitoreoAuditoria.objects.create(
            actor=self.observador, observador=self.observador,
            accion=MonitoreoAuditoria.CREAR_ZONA, zona_nombre="Norte",
            detalle="crea la zona Norte",
        )
        MonitoreoAuditoria.objects.create(
            actor=self.observador, observador=self.observador,
            accion=MonitoreoAuditoria.AGREGAR_LISTA, zona_nombre="",
            detalle="agrega a Ana",
        )
        MonitoreoAuditoria.objects.create(
            actor=self.ajeno, observador=self.ajeno,
            accion=MonitoreoAuditoria.ELIMINAR_ZONA, zona_nombre="Sur",
            detalle="crea la zona Sur",
        )
        self.client.force_login(self.observador)

    def _texto_del_fichero(self, respuesta):
        libro = load_workbook(BytesIO(respuesta.content))
        return " ".join(
            str(celda.value or "")
            for fila in libro.active.iter_rows()
            for celda in fila
        )

    def test_el_boton_de_excel_esta_en_su_auditoria(self):
        html = self.client.get(reverse("tickets:obs_mi_lista")).content.decode()
        self.assertIn(reverse("tickets:obs_auditoria_exportar"), html)
        # La descarga completa del administrador sigue sin colarse.
        self.assertNotIn(reverse("tickets:monitoreo_auditoria_exportar"), html)

    def test_descarga_solo_lo_que_es_suyo(self):
        respuesta = self.client.get(reverse("tickets:obs_auditoria_exportar"))
        self.assertEqual(respuesta.status_code, 200)
        self.assertIn("spreadsheetml", respuesta["Content-Type"])
        texto = self._texto_del_fichero(respuesta)
        self.assertIn("crea la zona Norte", texto)
        self.assertIn("agrega a Ana", texto)
        self.assertNotIn("crea la zona Sur", texto)

    def test_si_su_zona_no_permite_no_hay_boton_ni_descarga(self):
        self.zona.perm_exportar = False
        self.zona.save(update_fields=["perm_exportar"])
        html = self.client.get(reverse("tickets:obs_mi_lista")).content.decode()
        self.assertNotIn(reverse("tickets:obs_auditoria_exportar"), html)
        respuesta = self.client.get(reverse("tickets:obs_auditoria_exportar"))
        self.assertEqual(respuesta.status_code, 302)

    def test_sin_zonas_no_ofrece_el_boton(self):
        self.zona.delete()
        html = self.client.get(reverse("tickets:obs_mi_lista")).content.decode()
        self.assertNotIn(reverse("tickets:obs_auditoria_exportar"), html)

    def test_otro_rol_no_llega_a_la_descarga(self):
        admin = Usuario.objects.create_user(
            email="admin@x.com", password="x", nombre="Adm",
            rol=Usuario.Rol.ADMIN,
        )
        self.client.force_login(admin)
        self.client.raise_request_exception = False
        self.assertEqual(
            self.client.get(reverse("tickets:obs_auditoria_exportar")).status_code,
            403,
        )


@override_settings(STORAGES={
    "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
    "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
})
class ObservadorPapeleraDelPanelTest(TestCase):
    """La papelera del panel derecho quita de golpe las cuentas marcadas.

    Solo se puede quitar de la propia lista: el filtro `observador=...`
    está dentro de la vista, así que mandar a mano el id de una fila ajena
    no sirve de nada.
    """

    def setUp(self):
        self.observador = Usuario.objects.create_user(
            email="obs@x.com", password="x", nombre="Obs",
            rol=Usuario.Rol.OBSERVADOR,
        )
        self.otro = Usuario.objects.create_user(
            email="otro@x.com", password="x", nombre="Otro",
            rol=Usuario.Rol.OBSERVADOR,
        )
        self.uno = Usuario.objects.create_user(
            email="a@x.com", password="x", nombre="Ana", rol=Usuario.Rol.USUARIO,
        )
        self.dos = Usuario.objects.create_user(
            email="b@x.com", password="x", nombre="Beto", rol=Usuario.Rol.USUARIO,
        )
        self.tres = Usuario.objects.create_user(
            email="c@x.com", password="x", nombre="Caro", rol=Usuario.Rol.USUARIO,
        )
        self.zona = ZonaObservador.objects.create(
            observador=self.observador, nombre="Norte"
        )
        self.ajena = ZonaObservador.objects.create(
            observador=self.otro, nombre="Sur"
        )
        self.suyo = ElementoMiLista.objects.create(
            observador=self.observador, tipo=ElementoMiLista.TIPO_USUARIO,
            usuario=self.uno, usuario_email="a@x.com", zona=self.zona,
        )
        self.suyo_dos = ElementoMiLista.objects.create(
            observador=self.observador, tipo=ElementoMiLista.TIPO_USUARIO,
            usuario=self.dos, usuario_email="b@x.com", zona=self.zona,
        )
        self.de_ajena = ElementoMiLista.objects.create(
            observador=self.otro, tipo=ElementoMiLista.TIPO_USUARIO,
            usuario=self.tres, usuario_email="c@x.com", zona=self.ajena,
        )
        self.client.force_login(self.observador)

    def test_el_panel_ofrece_la_papelera(self):
        html = self.client.get(reverse("tickets:obs_mi_lista")).content.decode()
        self.assertIn('class="pan-papelera" id="panBorrar" disabled', html)
        self.assertIn('data-quitar-varios=', html)
        # Las filas llegan en su JSON con el id con el que se borran, que es
        # lo que la papelera recoge de cada casilla marcada.
        self.assertIn('"elemento_id": %d' % self.suyo.pk, html)
        self.assertIn('"elemento_id": %d' % self.suyo_dos.pk, html)

    def test_quitar_varios_bota_las_que_son_suyas(self):
        self.client.post(reverse("tickets:obs_quitar_varios"), {
            "elementos": [str(self.suyo.pk), str(self.de_ajena.pk)],
            "volver": reverse("tickets:obs_mi_lista"),
        })
        self.assertFalse(ElementoMiLista.objects.filter(pk=self.suyo.pk).exists())
        self.assertTrue(ElementoMiLista.objects.filter(pk=self.de_ajena.pk).exists())
        self.assertTrue(ElementoMiLista.objects.filter(pk=self.suyo_dos.pk).exists())

    def test_lo_que_no_sea_un_id_no_cae(self):
        self.client.post(reverse("tickets:obs_quitar_varios"), {
            "elementos": ["cosas", "9999999"],
        })
        self.assertEqual(
            ElementoMiLista.objects.filter(observador=self.observador).count(), 2
        )

    def test_sin_post_no_hace_nada(self):
        self.assertEqual(
            self.client.get(reverse("tickets:obs_quitar_varios")).status_code, 404
        )


@override_settings(STORAGES={
    "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
    "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
})
class ObservadorZonasConIconosTest(TestCase):
    """La página Zonas enseña cada carpeta con su icono y sus dos acciones:
    editar (el campo del nombre se despliega ahí mismo) y borrar."""

    def setUp(self):
        self.observador = Usuario.objects.create_user(
            email="obs@x.com", password="x", nombre="Obs",
            rol=Usuario.Rol.OBSERVADOR,
        )
        self.zona = ZonaObservador.objects.create(
            observador=self.observador, nombre="Norte"
        )
        self.client.force_login(self.observador)

    def test_cada_zona_lleva_icono_y_sus_acciones(self):
        html = self.client.get(reverse("tickets:obs_zonas")).content.decode()
        self.assertIn('class="obs-tarjeta-ico"', html)
        self.assertIn("#i-map-pin", html)
        self.assertIn(f'data-renombre="rn-{self.zona.pk}"', html)
        self.assertIn("#i-pencil", html)
        self.assertIn("#i-trash", html)
        self.assertIn(f'id="rn-{self.zona.pk}"', html)
        self.assertIn(
            'action="%s"' % reverse("tickets:obs_zona_eliminar",
                                    args=[self.zona.pk]),
            html,
        )

    def test_el_campo_de_renombrar_empieza_oculto(self):
        """La edición se despliega desde la tarjeta, no manda a otra página
        ni deja el campo a la vista estorbando."""
        html = self.client.get(reverse("tickets:obs_zonas")).content.decode()
        i_form = html.index('id="rn-%s"' % self.zona.pk)
        linea = html[i_form:i_form + 200]
        self.assertIn("hidden", linea)


@override_settings(STORAGES={
    "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
    "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
})
class ZonasConAccionesEnLasTresVistasTest(TestCase):
    """Las tres vistas de zonas dicen lo mismo: icono de carpeta y las
    acciones Ver, Editar (que despliega el nombre en la propia tarjeta) y
    Borrar con su papelera.

    Antes esas acciones solo estaban en la página Zonas: desde Monitoreo y
    desde el panel de Mi lista no había manera de quitar una carpeta sin
    salir de donde se estaba.
    """

    def setUp(self):
        self.observador = Usuario.objects.create_user(
            email="obs@x.com", password="x", nombre="Obs",
            rol=Usuario.Rol.OBSERVADOR,
        )
        self.zona = ZonaObservador.objects.create(
            observador=self.observador, nombre="Norte"
        )
        self.client.force_login(self.observador)

    def _acciones(self, html):
        """Comprueba la tira de acciones sobre la tarjeta de la zona."""
        for marca in (
            'data-renombre="rn-%d"' % self.zona.pk,
            'id="rn-%d" hidden' % self.zona.pk,
            "#i-eye",
            "#i-pencil",
            "#i-trash",
            'action="%s"' % reverse("tickets:obs_zona_eliminar",
                                    args=[self.zona.pk]),
            'name="zona_id" value="%d"' % self.zona.pk,
        ):
            with self.subTest(marca=marca):
                self.assertIn(marca, html)
        # El campo de renombrar va a la página de Zonas, que es la que sabe
        # corregir el nombre, y vuelve para acá con `volver`.
        self.assertIn('action="%s"' % reverse("tickets:obs_zonas"), html)

    def test_el_monitoreo_ofrece_ver_editar_y_borrar(self):
        respuesta = self.client.get(reverse("tickets:obs_monitoreo"))
        html = respuesta.content.decode()
        self._acciones(html)
        self.assertIn('name="volver" value="/observador/monitoreo/"', html)
        # El icono de la zona sigue mandando la tarjeta.
        self.assertIn('class="obs-zona-ico"', html)
        self.assertIn("#i-map-pin", html)

    def test_el_panel_de_mi_lista_también(self):
        html = self.client.get(reverse("tickets:obs_mi_lista")).content.decode()
        self._acciones(html)
        self.assertIn('name="volver" value="/observador/mi-lista/"', html)
        self.assertIn("panCerrarX", html)

    def test_borrar_desde_el_monitoreo_vuelve_al_monitoreo(self):
        ElementoMiLista.objects.create(
            observador=self.observador, tipo=ElementoMiLista.TIPO_USUARIO,
            usuario_email="ana@x.com", zona=self.zona,
        )
        respuesta = self.client.post(
            reverse("tickets:obs_zona_eliminar", args=[self.zona.pk]),
            {"volver": "/observador/monitoreo/"},
        )
        self.assertEqual(respuesta.status_code, 302)
        self.assertEqual(respuesta["Location"], "/observador/monitoreo/")
        self.assertFalse(
            ZonaObservador.objects.filter(pk=self.zona.pk).exists()
        )
        # La gente no se borra con la zona: se queda en la lista, sin carpeta.
        suelta = ElementoMiLista.objects.get(
            observador=self.observador, usuario_email="ana@x.com"
        )
        self.assertIsNone(suelta.zona)

    def test_renombrar_desde_el_panel_vuelve_al_panel(self):
        ElementoMiLista.objects.create(
            observador=self.observador, tipo=ElementoMiLista.TIPO_USUARIO,
            usuario_email="ana@x.com", zona=self.zona,
        )
        respuesta = self.client.post(reverse("tickets:obs_zonas"), {
            "zona_id": str(self.zona.pk),
            "nuevo_nombre": "Sede norte",
            "volver": "/observador/mi-lista/",
        })
        self.assertEqual(respuesta.status_code, 302)
        self.assertEqual(respuesta["Location"], "/observador/mi-lista/")
        self.zona.refresh_from_db()
        self.assertEqual(self.zona.nombre, "Sede norte")
        self.assertEqual(
            ElementoMiLista.objects.filter(
                observador=self.observador, zona=self.zona
            ).count(),
            1,
        )

    def test_el_panel_de_derecha_se_puede_retirar_de_verdad(self):
        """Cerrar no vacía la lista y se queda ahí: retira el panel y al
        elegir una tarjeta vuelve. Y la papelera usa la ventana de
        confirmación del aplicativo, no la del navegador."""
        html = self.client.get(reverse("tickets:obs_mi_lista")).content.decode()
        for marca in (
            'id="panCerrar"',
            'id="panCerrarX"',
            'grid.classList.add("is-cerrado")',
            'grid.classList.remove("is-cerrado")',
            "DoggerConfirma",
            'data-quitar-varios="%s"' % reverse("tickets:obs_quitar_varios"),
        ):
            with self.subTest(marca=marca):
                self.assertIn(marca, html)
        from django.contrib.staticfiles import finders
        with open(finders.find("css/dogger-theme.css"), encoding="utf-8") as f:
            css = f.read()
        self.assertIn(".pan-grid.is-cerrado", css)
        self.assertIn(".pan-right[hidden]", css)

    def test_sin_volver_sigue_landed_en_las_zonas(self):
        """La página Zonas no manda `volver` y se queda donde estaba."""
        respuesta = self.client.post(reverse("tickets:obs_zonas"), {
            "zona_id": str(self.zona.pk),
            "nuevo_nombre": "Sur",
        })
        self.assertEqual(respuesta.status_code, 302)
        self.assertEqual(respuesta["Location"], reverse("tickets:obs_zonas"))


class PaletaDeColorDeLasCarpetasTest(TestCase):
    """Cada tarjeta de carpeta lleva una gota en su esquina: abre la
    paleta con los diez colores cerrados y pinta la tarjeta con el que se
    elija. Lo que no venga de la paleta no se guarda."""

    def setUp(self):
        self.observador = Usuario.objects.create_user(
            email="obs@x.com", password="x", nombre="Obs",
            rol=Usuario.Rol.OBSERVADOR,
        )
        self.zona = ZonaObservador.objects.create(
            observador=self.observador, nombre="Norte"
        )
        self.client.force_login(self.observador)

    def test_las_tres_vistas_ofrecen_la_gota_y_la_paleta(self):
        for vista in ("obs_monitoreo", "obs_zonas", "obs_mi_lista"):
            with self.subTest(vista=vista):
                html = self.client.get(reverse("tickets:" + vista)).content.decode()
                self.assertIn('data-paleta="pl-%d"' % self.zona.pk, html)
                self.assertIn('id="pl-%d" hidden' % self.zona.pk, html)
                self.assertIn("#i-droplet", html)
                self.assertIn(
                    'action="%s"' % reverse("tickets:obs_zona_color"), html
                )

    def test_el_color_elegido_se_guarda_y_pinta(self):
        respuesta = self.client.post(reverse("tickets:obs_zona_color"), {
            "zona_id": str(self.zona.pk),
            "color": "#D62B1F",
            "volver": reverse("tickets:obs_zonas"),
        })
        self.assertEqual(respuesta.status_code, 302)
        self.assertEqual(respuesta["Location"], reverse("tickets:obs_zonas"))
        self.zona.refresh_from_db()
        self.assertEqual(self.zona.color, "#D62B1F")
        html = self.client.get(reverse("tickets:obs_zonas")).content.decode()
        self.assertIn("--zc: #D62B1F", html)
        # El marcado sigue siendo de la misma zona.
        self.assertIn('data-paleta="pl-%d"' % self.zona.pk, html)

    def test_un_color_de_fuera_de_la_paleta_no_se_guarda(self):
        r = self.client.post(reverse("tickets:obs_zona_color"), {
            "zona_id": str(self.zona.pk), "color": "javascript:alert(1)",
        })
        self.assertEqual(r.status_code, 404)
        self.zona.refresh_from_db()
        self.assertEqual(self.zona.color, "")
        # Tampoco se puede pintar la carpeta de otro observador.
        ajena = ZonaObservador.objects.create(
            observador=Usuario.objects.create_user(
                email="otro@x.com", password="x", nombre="Otro",
                rol=Usuario.Rol.OBSERVADOR,
            ),
            nombre="Ajena",
        )
        r = self.client.post(reverse("tickets:obs_zona_color"), {
            "zona_id": str(ajena.pk), "color": "#1D4ED8",
        })
        self.assertEqual(r.status_code, 404)
        ajena.refresh_from_db()
        self.assertEqual(ajena.color, "")

    def test_sin_color_vuelve_a_la_del_estado(self):
        self.zona.color = "#1D4ED8"
        self.zona.save(update_fields=["color"])
        self.client.post(reverse("tickets:obs_zona_color"), {
            "zona_id": str(self.zona.pk), "color": "",
        })
        self.zona.refresh_from_db()
        self.assertEqual(self.zona.color, "")
        html = self.client.get(reverse("tickets:obs_zonas")).content.decode()
        self.assertNotIn("--zc:", html)
