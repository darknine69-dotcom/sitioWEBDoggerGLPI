"""Exportaciones a Excel: estilo corporativo y gráficas de trazabilidad."""

from io import BytesIO

from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from apps.accounts.models import Usuario
from apps.tickets.models import Ticket


@override_settings(STORAGES={
    "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
    "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
})
class ExportesCorporativosTest(TestCase):
    def setUp(self):
        self.admin = Usuario.objects.create_user(
            email="adm@x.com", password="x", nombre="Adm", rol=Usuario.Rol.ADMIN
        )
        self.tec = Usuario.objects.create_user(
            email="tec@x.com", password="x", nombre="Tec", rol=Usuario.Rol.TECNICO
        )
        Ticket.objects.create(
            titulo="Impresora", descripcion="d", solicitante_nombre="Ana",
            solicitante_email="ana@x.com", tecnico_asignado=self.tec,
        )
        Ticket.objects.create(
            titulo="Red", descripcion="d", solicitante_nombre="Luis",
            solicitante_email="luis@x.com", estado=Ticket.Estado.RESUELTO,
        )
        self.client.force_login(self.admin)

    def test_reporte_con_graficas_y_estilo(self):
        respuesta = self.client.get(reverse("tickets:exportar_reportes"))
        self.assertEqual(respuesta.status_code, 200)
        self.assertIn("spreadsheetml.sheet", respuesta["Content-Type"])
        wb = load_workbook(BytesIO(respuesta.content))
        self.assertEqual(
            set(wb.sheetnames),
            {"Resumen", "Por técnico", "Por usuario", "Gráficas", "Tickets filtrados"},
        )
        # Encabezado corporativo: banda negra con texto blanco.
        resumen = wb["Resumen"]
        self.assertEqual(resumen["A1"].fill.start_color.rgb, "00171412")
        self.assertEqual(resumen["A1"].font.color.rgb, "00FFFFFF")
        # Hay gráficas nativas editables para la trazabilidad.
        con_graficas = [ws for ws in wb.worksheets if getattr(ws, "_charts", [])]
        self.assertTrue(con_graficas, "El reporte debe incluir gráficas")
        nombres = {type(ch).__name__ for ws in con_graficas for ch in ws._charts}
        self.assertIn("BarChart", nombres)
        self.assertIn("PieChart", nombres)

    def test_tickets_con_resumen_por_estado(self):
        respuesta = self.client.get(reverse("tickets:exportar_admin"))
        self.assertEqual(respuesta.status_code, 200)
        wb = load_workbook(BytesIO(respuesta.content))
        ws = wb["Tickets"]
        textos = [ws.cell(row=r, column=14).value for r in range(4, 10)]
        self.assertIn("Resumen por estado", textos)
        self.assertTrue(getattr(ws, "_charts", []))
